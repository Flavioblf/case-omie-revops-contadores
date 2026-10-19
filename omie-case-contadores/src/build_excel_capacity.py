"""Gera output/Desafio3_Capacity.xlsx: gap vs meta, framework de priorização e EC adicional — em fórmulas Excel.

A priorização usa o tier do Desafio 1c (matriz Fit x Engajamento, sql/01c_tiers_prioridade.sql).
Uso: python src/build_excel_capacity.py
"""
import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.formatting.rule import CellIsRule, ColorScaleRule

from load import ROOT, connect
from xlsx_style import (BORDA, BRL, F_ALERTA, F_SUB, F_TITULO, FILL_ALERTA, FILL_DESTAQUE, FILL_INPUT, PCT,
                        PCT1, WRAP, header, larguras, py)

OUT = ROOT / "output" / "Desafio3_Capacity.xlsx"
SQL_TIERS = ROOT / "sql" / "01c_tiers_prioridade.sql"

# ---------------------------------------------------------------- dados
con = connect()
con.execute(SQL_TIERS.read_text(encoding="utf-8").split("-- @name:")[0])  # (re)cria vw_tiers_prospects
franquias = con.execute("""
    SELECT id_franquia, nome_franquia, estado, regiao, meses_operando, num_executivos, meta_nmrr_mes, realizado_nmrr_mes,
           data_inicio
    FROM franquias ORDER BY id_franquia""").df()
execucao = con.execute("""
    SELECT id_executivo, id_franquia, estado, mes_ref, leads_trabalhados, reunioes_realizadas, propostas_enviadas,
           fechamentos, nmrr_gerado
    FROM execucao_comercial ORDER BY id_franquia, mes_ref""").df()
prospects = con.execute("""
    -- universo = mesma regra do Desafio 1a/1c (não é 'Cliente Ativo' e ainda não converteu) = view do tier 1c
    SELECT id_contador, estado, cidade, regiao, porte, num_clientes_ativos, erp_atual, status_canonico,
           estagio_aberto, estagio_aberto_ordem, tem_lead_parado,
           CASE WHEN num_clientes_ativos >= 50 AND num_funcionarios >= 6
                     AND erp_atual NOT IN ('Nenhum', 'Planilha', 'ContaAzul') THEN 1 ELSE 0 END AS icp,
           score_prioridade, tier_prioridade, fit_carteira, engajamento
    FROM vw_tiers_prospects
    ORDER BY id_contador""").df()
# propensão do Desafio 2: média das taxas de conversão das categorias do contador nos 5 drivers (600 contadores)
propensao = con.execute("""
    WITH b AS (
        SELECT id_contador, converteu_omie, porte, erp_atual,
               CASE WHEN anos_mercado < 8 THEN 1 WHEN anos_mercado < 16 THEN 2 ELSE 3 END AS f_anos,
               CASE WHEN num_clientes_ativos < 20 THEN 1 WHEN num_clientes_ativos < 50 THEN 2
                    WHEN num_clientes_ativos < 150 THEN 3 ELSE 4 END AS f_cli,
               CASE WHEN num_funcionarios <= 2 THEN 1 WHEN num_funcionarios <= 5 THEN 2
                    WHEN num_funcionarios <= 15 THEN 3 ELSE 4 END AS f_func
        FROM mercado_contadores)
    SELECT id_contador,
           (AVG(converteu_omie) OVER (PARTITION BY porte) + AVG(converteu_omie) OVER (PARTITION BY f_func)
            + AVG(converteu_omie) OVER (PARTITION BY f_cli) + AVG(converteu_omie) OVER (PARTITION BY erp_atual)
            + AVG(converteu_omie) OVER (PARTITION BY f_anos)) / 5 AS propensao
    FROM b""").df().set_index("id_contador").propensao
ticket_porte = con.execute("""
    SELECT m.porte, ROUND(AVG(c.mrr_inicial), 0) AS ticket
    FROM clientes_ativos c JOIN mercado_contadores m USING (id_contador)
    GROUP BY m.porte ORDER BY ticket DESC""").df()
leads_por_ec = float(con.execute("SELECT AVG(leads_trabalhados) FROM execucao_comercial").fetchone()[0])
churn = float(con.execute("SELECT AVG(CASE WHEN status_cliente = 'Churned' THEN 1.0 ELSE 0 END) FROM clientes_ativos").fetchone()[0])
estados_mercado = [r[0] for r in con.execute("SELECT DISTINCT estado FROM mercado_contadores ORDER BY 1").fetchall()]
con.close()

# território: a franquia do estado; SP é dividido por cidade na fórmula; estado sem franquia fica sem dono
estado_franquia = dict(zip(franquias.estado, franquias.id_franquia))
NF = len(franquias)
F0, F1 = 5, 4 + NF  # linhas das franquias na aba Gap_Meta
NE = len(execucao) + 1
NP = len(prospects) + 1

wb = Workbook()

# ---------------------------------------------------------------- LEIA-ME
ws = wb.active
ws.title = "LEIA-ME"
linhas = [
    ("Desafio 3 — Priorização Comercial: Capacity & Foco dos ECs", F_TITULO),
    ("Tudo calculado por fórmulas a partir da base do case. Células verdes = resultados principais.", None),
    ("", None),
    ("Abas", F_SUB),
    ("1. Premissas — só parâmetros tirados da base: capacidade por EC, churn, corte de lead parado, território e ticket.", None),
    ("2. Gap_Meta (3a) — meta vs realizado por franquia e por EC equivalente, funil vs rede e tendência.", None),
    ("3. Regioes (3a) — onde se concentra o gap e o risco.", None),
    ("4. Simulador_Funil — quanto vale cada ponto percentual de conversão em cada etapa do funil.", None),
    (f"5. Priorizacao (3b) — {len(prospects)} contadores endereçáveis (regra do 1a: não é 'Cliente Ativo' e converteu_omie = 0) "
     "com tier de prioridade (Desafio 1c), intenção (pipeline), MRR esperado, onda e ordem de ataque.", None),
    ("6. Ondas_Capacidade (3b) — o framework: quem entra em cada onda, o que o EC faz e por quê.", None),
    ("7. EC_Adicional (3c) — onde alocar 1 EC adicional: demanda por EC × conversão do funil.", None),
    ("8. Execucao — base execucao_comercial (apoio).", None),
    ("", None),
    ("Escolhas e premissas", F_SUB),
    ("• Meta e realizado: aba franquias (fonte oficial de meta). O funil (taxas, tendência) vem de execucao_comercial: "
     "usamos apenas proporções, pois o nMRR da execução é 10 a 27x menor que o oficial (franquias.realizado_nmrr_mes ÷ média mensal de execucao_comercial.nmrr_gerado, por franquia).", None),
    ("• Nível de EC: o id_executivo se repete dentro da franquia (1 ID para todos os ECs), então não há como seguir um EC individual. "
     "Usamos o EC equivalente (franquia ÷ nº de ECs) e a dispersão (% de EC-mês sem fechamento).", None),
    ("• Prioridade = Intenção (lead ativo no pipeline) > Tier do Desafio 1c > MRR esperado. A intenção vem "
     "primeiro porque é sinal real de compra; o perfil ainda precisa ser validado em piloto (Desafio 2, aba Como_Avaliar).", None),
    ("• Tier do 1c = matriz Fit x Engajamento: Fit = carteira (nº de clientes, driver do Desafio 2), "
     "Engajamento = estágio do lead aberto nos últimos 90 dias. Sem pesos arbitrários.", None),
    ("• MRR esperado = ticket médio do porte × (1 − churn): o valor de cada contador, sem nota inventada.", None),
    ("• Território: cada contador fica com a franquia do seu estado. SP: cidade de São Paulo = SP Capital, demais cidades = "
     "SP Interior (pelo nome das franquias). MT, MA, RO e MS não têm franquia: ficam como 'Sem franquia', sem dono atribuído.", None),
    ("• Sem premissas externas: nenhum custo, ramp-up, SLA ou corte que não venha da base.", None),
]
for i, (txt, fonte) in enumerate(linhas, start=1):
    c = ws.cell(row=i, column=1, value=txt)
    c.alignment = WRAP
    if fonte:
        c.font = fonte
ws.column_dimensions["A"].width = 150

# ---------------------------------------------------------------- Premissas
ws = wb.create_sheet("Premissas")
ws["A1"], ws["A1"].font = "Parâmetros (todos tirados da base do case)", F_TITULO
header(ws, 3, ["Parâmetro", "Valor", "Fonte na base"])
prem = [
    ("Leads trabalhados por EC/mês (capacidade)", round(leads_por_ec, 1), "Média de leads_trabalhados (execucao_comercial)", "0.0"),
    ("Churn dos contratos", round(churn, 4), "Contratos Churned ÷ total (clientes_ativos)", PCT1),
    ("Lead parado a partir de (dias)", 90, "Nenhum lead Ganho da base fechou com mais de 89 dias (leads_pipeline)", "0"),
]
for k, (nome, v, fonte, fmt) in enumerate(prem):
    r = 4 + k
    ws.cell(row=r, column=1, value=nome)
    ws.cell(row=r, column=2, value=v).number_format = fmt
    ws.cell(row=r, column=3, value=fonte)
# B4 capacidade | B5 churn | B6 corte de lead parado
ws["A9"], ws["A9"].font = "Território (estado → franquia)", F_SUB
header(ws, 10, ["Estado", "Franquia", "Cobertura"])
for k, uf in enumerate(estados_mercado):
    r = 11 + k
    ws.cell(row=r, column=1, value=uf)
    if uf == "SP":
        ws.cell(row=r, column=2, value="FRQ001 / FRQ002")
        ws.cell(row=r, column=3, value="Própria (cidade de São Paulo = SP Capital; demais = SP Interior)")
    elif uf in estado_franquia:
        ws.cell(row=r, column=2, value=estado_franquia[uf])
        ws.cell(row=r, column=3, value="Própria")
    else:
        ws.cell(row=r, column=2, value="Sem franquia")
        ws.cell(row=r, column=3, value="Sem franquia no estado (sem cobertura)")
MAP0, MAP1 = 11, 10 + len(estados_mercado)
r0 = MAP1 + 3
ws.cell(row=r0 - 1, column=1, value="Ticket MRR por porte (média de mrr_inicial em clientes_ativos)").font = F_SUB
header(ws, r0, ["Porte", "Ticket (R$)"])
for k, row in enumerate(ticket_porte.itertuples(index=False)):
    ws.cell(row=r0 + 1 + k, column=1, value=row.porte)
    ws.cell(row=r0 + 1 + k, column=2, value=py(row.ticket)).number_format = BRL
TK0, TK1 = r0 + 1, r0 + len(ticket_porte)
larguras(ws, {"A": 52, "B": 20, "C": 70})
MAPA = f"Premissas!$A${MAP0}:$B${MAP1}"
TICKET = f"Premissas!$A${TK0}:$B${TK1}"

# ---------------------------------------------------------------- Execucao (apoio)
ws = wb.create_sheet("Execucao")
header(ws, 1, ["id_executivo", "id_franquia", "estado", "mes_ref", "leads_trabalhados", "reunioes_realizadas",
               "propostas_enviadas", "fechamentos", "nmrr_gerado", "semestre", "ec_mes_sem_fechamento"])
for i, row in enumerate(execucao.itertuples(index=False), start=2):
    for j, v in enumerate(row, start=1):
        ws.cell(row=i, column=j, value=py(v))
    ws[f"J{i}"] = f'=IF(VALUE(RIGHT(D{i},2))<=6,"H1","H2")'
    ws[f"K{i}"] = f"=IF(H{i}=0,1,0)"
ws.freeze_panes = "A2"
ws.auto_filter.ref = f"A1:K{NE}"


def ex(col):
    return f"Execucao!${col}$2:${col}${NE}"


# ---------------------------------------------------------------- Gap_Meta (3a)
ws = wb.create_sheet("Gap_Meta")
ws["A1"], ws["A1"].font = "3a — Franquias vs meta de nMRR: onde está o gap e por quê", F_TITULO
ws["A2"] = ("Meta/realizado: aba franquias (oficial). Funil e tendência: execucao_comercial 2024 (proporções). "
            "Índice vs benchmark = taxa da franquia ÷ taxa da rede (1,00 = igual à rede).")
cols = ["id", "Franquia", "UF", "Região", "Meses operando", "Nº ECs", "Meta nMRR/mês", "Realizado nMRR/mês",
        "Atingimento", "Gap (R$)", "Meta por EC", "Realizado por EC", "Realizado por EC ÷ média da rede",
        "Leads", "Reuniões", "Propostas", "Fechamentos",
        "Lead→Reunião", "Reunião→Proposta", "Proposta→Fechamento",
        "Índice L→R", "Índice R→P", "Índice P→F", "Etapa mais fraca vs rede",
        "Tendência H2 vs H1", "% EC-mês sem fechamento", "Situação", "Início (data_inicio)"]
header(ws, 4, cols)
for k, row in enumerate(franquias.itertuples(index=False)):
    r = F0 + k
    for j, v in enumerate([row.id_franquia, row.nome_franquia, row.estado, row.regiao, row.meses_operando,
                           row.num_executivos, row.meta_nmrr_mes, row.realizado_nmrr_mes], start=1):
        ws.cell(row=r, column=j, value=py(v))
    f = {
        "I": f"=H{r}/G{r}", "J": f"=MAX(0,G{r}-H{r})", "K": f"=G{r}/F{r}", "L": f"=H{r}/F{r}",
        "M": f"=L{r}/L${F1 + 1}",
        "N": f'=SUMIFS({ex("E")},{ex("B")},A{r})', "O": f'=SUMIFS({ex("F")},{ex("B")},A{r})',
        "P": f'=SUMIFS({ex("G")},{ex("B")},A{r})', "Q": f'=SUMIFS({ex("H")},{ex("B")},A{r})',
        "R": f"=O{r}/N{r}", "S": f"=P{r}/O{r}", "T": f"=Q{r}/P{r}",
        "U": f"=R{r}/R${F1 + 2}", "V": f"=S{r}/S${F1 + 2}", "W": f"=T{r}/T${F1 + 2}",
        "X": (f'=IF(U{r}=MIN(U{r}:W{r}),"Lead→Reunião",IF(V{r}=MIN(U{r}:W{r}),"Reunião→Proposta","Proposta→Fechamento"))'),
        "Y": (f'=SUMIFS({ex("I")},{ex("B")},A{r},{ex("J")},"H2")/SUMIFS({ex("I")},{ex("B")},A{r},{ex("J")},"H1")-1'),
        "Z": f'=AVERAGEIFS({ex("K")},{ex("B")},A{r})',
        "AA": f'=IF(I{r}>=1,"Na meta","Abaixo da meta")',
    }
    for col, formula in f.items():
        ws[f"{col}{r}"] = formula
    ws[f"AB{r}"] = pd.Timestamp(row.data_inicio).to_pydatetime()
    ws[f"AB{r}"].number_format = "dd/mm/yyyy"
    for col in "GHJKL":
        ws[f"{col}{r}"].number_format = BRL
    for col in "IRST":
        ws[f"{col}{r}"].number_format = PCT1
    for col in "MUVW":
        ws[f"{col}{r}"].number_format = "0.00"
    ws[f"Y{r}"].number_format = "+0%;-0%;0%"
    ws[f"Z{r}"].number_format = PCT
tr = F1 + 1  # total
br = F1 + 2  # benchmark
ws.cell(row=tr, column=2, value="Total da rede").font = F_SUB
for col in "FGHJNOPQ":
    ws[f"{col}{tr}"] = f"=SUM({col}{F0}:{col}{F1})"
ws[f"I{tr}"] = f"=H{tr}/G{tr}"
ws[f"K{tr}"] = f"=G{tr}/F{tr}"
ws[f"L{tr}"] = f"=H{tr}/F{tr}"
ws.cell(row=br, column=2, value="Benchmark da rede (funil)").font = F_SUB
ws[f"R{br}"], ws[f"S{br}"], ws[f"T{br}"] = f"=O{tr}/N{tr}", f"=P{tr}/O{tr}", f"=Q{tr}/P{tr}"
ws[f"Z{br}"] = f'=AVERAGE({ex("K")})'
for col in "GHJKL":
    ws[f"{col}{tr}"].number_format = BRL
for a in (f"I{tr}", f"R{br}", f"S{br}", f"T{br}"):
    ws[a].number_format = PCT1
ws[f"Z{br}"].number_format = PCT
ws[f"J{tr}"].fill = FILL_DESTAQUE
ws.conditional_formatting.add(f"I{F0}:I{F1}", CellIsRule(operator="lessThan", formula=["1"], fill=FILL_ALERTA))
ws.conditional_formatting.add(f"U{F0}:W{F1}", ColorScaleRule(start_type="num", start_value=0.8, start_color="F8696B",
                                                             mid_type="num", mid_value=1, mid_color="FFFFFF",
                                                             end_type="num", end_value=1.2, end_color="63BE7B"))

# Leitura com números vivos
s = F1 + 4
ws.cell(row=s, column=1, value="Leitura").font = F_SUB
indic = [
    ("Gap total das franquias abaixo da meta (R$/mês)", f"=J{tr}", BRL),
    ("Franquias abaixo da meta", f'=COUNTIF(I{F0}:I{F1},"<1")', "0"),
    ("Correlação meses_operando × atingimento", f"=CORREL(E{F0}:E{F1},I{F0}:I{F1})", "0.00"),
    ("Correlação nº de ECs × atingimento", f"=CORREL(F{F0}:F{F1},I{F0}:I{F1})", "0.00"),
    ("Faixa dos índices L→R e R→P entre franquias", f'=FIXED(MIN(U{F0}:V{F1}),2)&" a "&FIXED(MAX(U{F0}:V{F1}),2)', None),
    ("Faixa do índice P→F entre franquias", f'=FIXED(MIN(W{F0}:W{F1}),2)&" a "&FIXED(MAX(W{F0}:W{F1}),2)', None),
    ("Correlação índice P→F × atingimento", f"=CORREL(W{F0}:W{F1},I{F0}:I{F1})", "0.00"),
    ("Correlação antiguidade pela data_inicio × atingimento (início mais cedo = mais antiga)",
     f"=-CORREL(AB{F0}:AB{F1},I{F0}:I{F1})", "0.00"),
    ("Correlação meses_operando × nº de ECs", f"=CORREL(E{F0}:E{F1},F{F0}:F{F1})", "0.00"),
]
for k, (nome, f, fmt) in enumerate(indic):
    ws.cell(row=s + 1 + k, column=1, value=nome)
    c = ws.cell(row=s + 1 + k, column=4, value=f)
    if fmt:
        c.number_format = fmt
    c.fill = FILL_DESTAQUE
conclusoes = [
    "1. O funil não explica o gap: L→R e R→P são quase iguais entre franquias; P→F varia mais, mas quase não se relaciona "
    "com o atingimento (correlação ~0,15, contra ~0,64 do nº de ECs). Leads trabalhados por EC variam pouco (15 a 19 por mês).",
    "2. O gap está nas franquias pequenas: das 7 franquias com 1 a 2 ECs, 6 estão abaixo da meta (só o ES bate). O CE tem a "
    "maior meta por EC da rede (R$ 10 mil) e realiza acima da média por EC: parte do gap é meta, não execução.",
    "3. Governança: meses_operando conta a mesma história (correlação ~0,99 com o nº de ECs), mas contradiz a data_inicio "
    "(coluna AB): pela data de início, as franquias mais antigas atingem MENOS (correlação ~-0,37). Por isso a leitura usa o "
    "nº de ECs, que não tem contradição na base.",
    "4. RJ: 5 ECs, maior gap absoluto (R$ 4 mil) e realizado por EC de R$ 5,6 mil, abaixo da média da rede.",
    "5. Maiores quedas de nMRR na execução do 2º semestre: DF, MG, SP Interior e SC (coluna Y) — inclusive em franquias na meta.",
    "Nível de EC: id_executivo não identifica o EC; usamos o EC equivalente (franquia ÷ nº de ECs) e a dispersão "
    "(% de EC-mês sem fechamento). Recomendação: ID único de EC no CRM.",
]
for k, t in enumerate(conclusoes):
    ws.cell(row=s + 11 + k, column=1, value=t)
ws.freeze_panes = "C5"
ws.row_dimensions[4].height = 48
larguras(ws, {"A": 9, "B": 21, "C": 5, "D": 12, "E": 9, "F": 7, "G": 12, "H": 12, "I": 11, "J": 10, "K": 11,
              "L": 11, "M": 12, "N": 7, "O": 8, "P": 9, "Q": 11, "R": 10, "S": 10, "T": 11, "U": 8, "V": 8,
              "W": 8, "X": 22, "Y": 11, "Z": 11, "AA": 16, "AB": 12})
GAP = "Gap_Meta"

# ---------------------------------------------------------------- Regioes (3a)
ws = wb.create_sheet("Regioes")
ws["A1"], ws["A1"].font = "3a — Concentração do gap e do risco por região", F_TITULO
header(ws, 3, ["Região", "Franquias", "Abaixo da meta", "% abaixo", "ECs", "Meta nMRR/mês", "Realizado nMRR/mês",
               "Atingimento da região", "Gap das franquias abaixo da meta (R$)", "% do gap total", "Meses operando (média)"])
gr = lambda c: f"{GAP}!${c}${F0}:${c}${F1}"  # noqa: E731
regioes = ["Sudeste", "Sul", "Centro-Oeste", "Nordeste", "Norte"]
for k, reg in enumerate(regioes):
    r = 4 + k
    ws.cell(row=r, column=1, value=reg)
    ws.cell(row=r, column=2, value=f'=COUNTIF({gr("D")},A{r})')
    ws.cell(row=r, column=3, value=f'=COUNTIFS({gr("D")},A{r},{gr("I")},"<1")')
    ws.cell(row=r, column=4, value=f"=C{r}/B{r}").number_format = PCT
    ws.cell(row=r, column=5, value=f'=SUMIF({gr("D")},A{r},{gr("F")})')
    ws.cell(row=r, column=6, value=f'=SUMIF({gr("D")},A{r},{gr("G")})').number_format = BRL
    ws.cell(row=r, column=7, value=f'=SUMIF({gr("D")},A{r},{gr("H")})').number_format = BRL
    ws.cell(row=r, column=8, value=f"=G{r}/F{r}").number_format = PCT1
    ws.cell(row=r, column=9, value=f'=SUMIF({gr("D")},A{r},{gr("J")})').number_format = BRL
    ws.cell(row=r, column=10, value=f"=I{r}/SUM($I$4:$I$8)").number_format = PCT
    ws.cell(row=r, column=11, value=f'=AVERAGEIF({gr("D")},A{r},{gr("E")})').number_format = "0.0"
ws["A9"] = "Total"
ws["A9"].font = F_SUB
for col in "BCEFGI":
    ws[f"{col}9"] = f"=SUM({col}4:{col}8)"
    ws[f"{col}9"].number_format = BRL if col in "FGI" else "0"
ws["H9"] = "=G9/F9"
ws["H9"].number_format = PCT1
ws.conditional_formatting.add("J4:J8", ColorScaleRule(start_type="min", start_color="FFFFFF",
                                                     end_type="max", end_color="F8696B"))
ws["A11"], ws["A11"].font = "Leitura", F_SUB
ws["A12"] = ("• Norte, Nordeste e Centro-Oeste: 100% das franquias abaixo da meta e a região inteira abaixo da meta "
             "(atingimento de 74%, 82% e 89%). São as regiões das menores franquias: 1 a 3 ECs (Norte e Nordeste, 1 a 2).")
ws["A13"] = ("• Sudeste e Sul batem a meta como região (107% e 108%): o gap é pontual, só do RJ (maior gap "
             "da rede) e do RS.")
ws["A14"] = ("• A rede como um todo bate a meta (R$ 299 mil realizados contra R$ 297,5 mil): os R$ 19,5 mil são o gap "
             "das 9 franquias abaixo da meta.")
chart = BarChart()
chart.title = "Gap de nMRR por região (R$/mês)"
chart.add_data(Reference(ws, min_col=9, min_row=3, max_row=8), titles_from_data=True)
chart.set_categories(Reference(ws, min_col=1, min_row=4, max_row=8))
chart.legend = None
chart.height, chart.width = 7, 14
ws.add_chart(chart, "A17")
larguras(ws, {"A": 14, "B": 10, "C": 10, "D": 9, "E": 6, "F": 14, "G": 14, "H": 11, "I": 11, "J": 11, "K": 12})
ws.row_dimensions[3].height = 32

# ---------------------------------------------------------------- Simulador de funil
ws = wb.create_sheet("Simulador_Funil")
ws["A1"], ws["A1"].font = "Quanto vale +1 p.p. em cada etapa do funil?", F_TITULO
ws["A2"] = ("Fechamentos por lead = (L→R) × (R→P) × (P→F). Um ganho de conversão aumenta o nMRR na mesma proporção "
            "(mesmo volume de leads e ticket). Resultado em % é independente da escala de receita.")
header(ws, 4, ["Etapa", "Taxa atual (rede)", "Ganho (p.p.)", "Taxa simulada"])
etapas = [("Lead → Reunião", "R"), ("Reunião → Proposta", "S"), ("Proposta → Fechamento", "T")]
for k, (nome, col) in enumerate(etapas):
    r = 5 + k
    ws.cell(row=r, column=1, value=nome)
    ws.cell(row=r, column=2, value=f"={GAP}!{col}{br}").number_format = PCT1
    c = ws.cell(row=r, column=3, value=0.05 if col == "T" else 0)
    c.number_format, c.fill = '+0.0%;-0.0%;0.0%', FILL_INPUT
    ws.cell(row=r, column=4, value=f"=B{r}+C{r}").number_format = PCT1
res = [
    ("Fechamentos por lead — atual", "=B5*B6*B7", PCT1),
    ("Fechamentos por lead — simulado", "=D5*D6*D7", PCT1),
    ("Uplift no nMRR", "=B10/B9-1", "+0%"),
    ("nMRR adicional/mês — rede (escala oficial)", f"=B11*{GAP}!H{tr}", BRL),
    ("nMRR adicional/mês — 9 franquias abaixo da meta (escala oficial)",
     f'=B11*SUMIFS({GAP}!$H${F0}:$H${F1},{GAP}!$I${F0}:$I${F1},"<1")', BRL),
    ("Cobertura do gap de meta", f"=B13/{GAP}!J{tr}", PCT),
]
for k, (nome, f, fmt) in enumerate(res):
    r = 9 + k
    ws.cell(row=r, column=1, value=nome)
    c = ws.cell(row=r, column=2, value=f)
    c.number_format = fmt
    if r >= 11:
        c.fill = FILL_DESTAQUE
ws["A16"] = ("Leitura: a etapa Proposta → Fechamento é a de menor taxa (~19%) — cada p.p. ali vale mais que nas outras etapas. "
             "Alavancas: qualificação antes da proposta, proposta padronizada e follow-up.")
ws["A16"].alignment = WRAP
ws.merge_cells("A16:D18")
larguras(ws, {"A": 60, "B": 16, "C": 14, "D": 14})

# ---------------------------------------------------------------- Priorizacao (3b)
ws = wb.create_sheet("Priorizacao")
cols_p = ["id_contador", "UF", "Cidade", "Região", "Porte", "Nº clientes", "ERP", "Status pelos contratos (clientes_ativos)",
          "Estágio do lead ativo", "Ordem do estágio", "Lead parado (>90d)", "ICP",
          "Score Fit + Engajamento (1c, 1 a 5)", "Tier (1c)", "Franquia responsável", "Cobertura",
          "Ticket esperado (R$)", "MRR esperado (R$)", "Onda", "Ação da onda", "Ordem na franquia",
          "Fit da carteira (1c)", "Engajamento (1c)", "Propensão (Desafio 2)"]
header(ws, 1, cols_p)
campos = ["id_contador", "estado", "cidade", "regiao", "porte", "num_clientes_ativos", "erp_atual", "status_canonico",
          "estagio_aberto", "estagio_aberto_ordem", "tem_lead_parado", "icp", "score_prioridade", "tier_prioridade"]
campos_fim = ["fit_carteira", "engajamento"]
P = lambda c: f"${c}$2:${c}${NP}"  # noqa: E731
for i, row in enumerate(prospects[campos].itertuples(index=False), start=2):
    for j, v in enumerate(row, start=1):
        ws.cell(row=i, column=j, value=py(v))
    for j, v in enumerate(prospects.loc[i - 2, campos_fim], start=22):
        ws.cell(row=i, column=j, value=py(v))
    ws.cell(row=i, column=24, value=round(float(propensao[row[0]]), 4)).number_format = PCT
    f = {
        "O": (f'=IF(B{i}="SP",IF(C{i}="São Paulo","FRQ001","FRQ002"),VLOOKUP(B{i},{MAPA},2,FALSE))'),
        "P": f'=IF(O{i}="Sem franquia","Sem franquia","Própria")',
        "Q": f"=VLOOKUP(E{i},{TICKET},2,FALSE)",
        "R": f"=Q{i}*(1-Premissas!$B$5)",
        "S": (f'=IF(J{i}>=3,1,IF(N{i}="Alta",2,IF(K{i}=1,3,IF(N{i}="Média",4,5))))'),
        "T": (f'=CHOOSE(S{i},"1. Fechar (lead ativo em Reunião/Proposta)","2. Prospectar tier Alta",'
              f'"3. Requalificar ou encerrar lead parado","4. Nutrir (marketing/inbound)","5. Régua automatizada")'),
        "U": (f'=COUNTIFS({P("O")},O{i},{P("S")},"<"&S{i})+COUNTIFS({P("O")},O{i},{P("S")},S{i},{P("R")},">"&R{i})+1'),
    }
    for col, formula in f.items():
        ws[f"{col}{i}"] = formula
    ws[f"M{i}"].number_format = "0"
    ws[f"Q{i}"].number_format = BRL
    ws[f"R{i}"].number_format = BRL
ws.freeze_panes = "B2"
ws.auto_filter.ref = f"A1:X{NP}"
ws.row_dimensions[1].height = 45
larguras(ws, {"A": 11, "B": 5, "C": 15, "D": 11, "E": 17, "F": 9, "G": 10, "H": 19, "I": 16, "J": 8, "K": 8,
              "L": 5, "M": 11, "N": 11, "O": 10, "P": 10, "Q": 10, "R": 11, "S": 6, "T": 38, "U": 9,
              "V": 11, "W": 11, "X": 12})
PR = lambda c: f"Priorizacao!${c}$2:${c}${NP}"  # noqa: E731

# ---------------------------------------------------------------- Ondas_Capacidade (3b)
ws = wb.create_sheet("Ondas_Capacidade")
ws["A1"], ws["A1"].font = "3b — Framework de priorização: Intenção → Fit x Engajamento (1c) → Valor", F_TITULO
ws["A2"] = ("Ordem de ataque dentro de cada franquia: primeiro a onda (abaixo), depois o maior MRR esperado "
            "(ticket do porte × (1 − churn)). O Fit é a carteira, driver nº 3 do Desafio 2; entre os 303 prospects "
            "ela acompanha o score de propensão do D2 (correlação ~0,8, coluna X da Priorizacao). "
            "Lista completa e ordem na aba Priorizacao.")
header(ws, 4, ["Onda", "Quem entra", "O que o EC faz", "Por quê", "Critério de saída", "Contadores",
               "MRR esperado (R$)"])
ondas = [
    ("1. Fechar", "Lead aberto há até 90 dias em Reunião Agendada ou Proposta Enviada — qualquer tier",
     "Follow-up da proposta e negociação",
     "É o único sinal real de compra; nenhum Ganho da base fechou depois de 89 dias", "Ganho ou Perdido com motivo"),
    ("2. Prospectar tier Alta", "Tier Alta do Desafio 1c (Fit Alto + Morno) que não está na Onda 1",
     "Outbound do EC: diagnóstico e reunião",
     "Carteira acima de 24 clientes (driver do D2, acompanha a propensão) e lead em andamento", "Reunião ou descarte"),
    ("3. Requalificar parado", "Lead aberto há mais de 90 dias (fora das ondas 1 e 2)",
     "Uma tentativa de requalificação; sem avanço, encerrar como Perdido",
     "Lead com mais de 89 dias nunca virou Ganho na base: limpar libera tempo de EC",
     "Lead volta a ativo ou é encerrado"),
    ("4. Nutrir", "Tier Média fora das ondas 1 a 3", "Marketing (conteúdo, eventos, inbound); EC só se o contador levantar a mão",
     "Fit ou engajamento parcial: ainda não justifica tempo de EC", "Pedido de contato → Onda 2"),
    ("5. Régua automatizada", "Tier Baixa fora das ondas 1 a 3", "Comunicação automatizada, sem tempo de EC",
     "Sem fit e sem engajamento", "Resposta → requalificar"),
]
for k, (nome, crit, acao, sla, saida) in enumerate(ondas):
    r = 5 + k
    for j, v in enumerate([nome, crit, acao, sla, saida], start=1):
        ws.cell(row=r, column=j, value=v).alignment = WRAP
    ws.cell(row=r, column=6, value=f"=COUNTIF({PR('S')},{k + 1})")
    ws.cell(row=r, column=7, value=f"=SUMIF({PR('S')},{k + 1},{PR('R')})").number_format = BRL
    ws.row_dimensions[r].height = 45
ws["E10"], ws["F10"], ws["G10"] = "Total", "=SUM(F5:F9)", "=SUM(G5:G9)"
ws["G10"].number_format = BRL

ws["A13"], ws["A13"].font = "Carga das ondas 1-3 (as que consomem tempo de EC) vs capacidade mensal por franquia", F_SUB
header(ws, 14, ["id", "Franquia", "ECs", "Capacidade (leads/mês)", "Onda 1", "Onda 2", "Onda 3", "Ondas 1-3",
                "% da capacidade de 1 mês", "MRR esperado ondas 1-2 (R$)", "Onda 4", "Onda 5"])
for k in range(NF):
    r = 15 + k
    g = F0 + k
    ws.cell(row=r, column=1, value=f"={GAP}!A{g}")
    ws.cell(row=r, column=2, value=f"={GAP}!B{g}")
    ws.cell(row=r, column=3, value=f"={GAP}!F{g}")
    ws.cell(row=r, column=4, value=f"=C{r}*Premissas!$B$4").number_format = "0"
    for j, onda in enumerate([1, 2, 3]):
        ws.cell(row=r, column=5 + j, value=f"=COUNTIFS({PR('O')},$A{r},{PR('S')},{onda})")
    ws.cell(row=r, column=8, value=f"=SUM(E{r}:G{r})")
    ws.cell(row=r, column=9, value=f"=H{r}/D{r}").number_format = PCT
    ws.cell(row=r, column=10, value=(f"=SUMIFS({PR('R')},{PR('O')},$A{r},{PR('S')},1)"
                                     f"+SUMIFS({PR('R')},{PR('O')},$A{r},{PR('S')},2)")).number_format = BRL
    ws.cell(row=r, column=11, value=f"=COUNTIFS({PR('O')},$A{r},{PR('S')},4)")
    ws.cell(row=r, column=12, value=f"=COUNTIFS({PR('O')},$A{r},{PR('S')},5)")
OC1 = 14 + NF
ws.conditional_formatting.add(f"I15:I{OC1}", ColorScaleRule(start_type="min", start_color="FFFFFF",
                                                            end_type="max", end_color="F8696B"))
ws.cell(row=OC1 + 2, column=1, value="Leitura").font = F_SUB
leit = [
    "• Intenção primeiro: a Onda 1 é o único grupo com sinal real de compra (lead ativo em estágio avançado).",
    "• O tier Alta vem em seguida, mas o perfil precisa ser validado: rodar as Ondas 1-2 como piloto com grupo de controle (Desafio 2).",
    "• A Onda 3 libera capacidade: limpar leads parados devolve horas de EC e torna a cobertura de pipeline confiável.",
    "• Contadores em estados sem franquia (MT, MA, RO, MS) ficam fora da tabela por franquia: não há EC responsável.",
    "• A capacidade nominal (44 ECs × 16,5 leads/mês ≈ 726 leads) supera os 303 endereçáveis da base. "
    "O que importa é a ORDEM de ataque, não a falta de capacidade.",
]
for k, t in enumerate(leit):
    ws.cell(row=OC1 + 3 + k, column=1, value=t)
larguras(ws, {"A": 22, "B": 40, "C": 38, "D": 26, "E": 30, "F": 11, "G": 16, "H": 10, "I": 13, "J": 18,
              "K": 8, "L": 8})
ws.row_dimensions[14].height = 45

# ---------------------------------------------------------------- EC_Adicional (3c)
ws = wb.create_sheet("EC_Adicional")
ws["A1"], ws["A1"].font = "3c — Onde alocar 1 EC adicional?", F_TITULO
ws["A2"] = ("Critério (só dados da base): o EC vai para onde há mais fechamentos potenciais por EC = endereçáveis do "
            "território por EC × conversão lead → fechamento da própria franquia (execucao_comercial). "
            "Sensibilidade: ranking refeito só pela demanda e só pela conversão.")
header(ws, 4, ["id", "Franquia", "Meses de operação", "ECs", "Endereçáveis no território", "Endereçáveis por EC",
               "Conversão lead → fechamento", "Fechamentos potenciais por EC", "Gap (R$)", "Atingimento",
               "Realizado por EC", "Ranking (fechamentos potenciais)", "Ranking (demanda por EC)",
               "Ranking (conversão)", "Leads trabalhados por EC-mês (carga)"])
E1 = 4 + NF
for k in range(NF):
    r = 5 + k
    g = F0 + k
    ws.cell(row=r, column=1, value=f"={GAP}!A{g}")
    ws.cell(row=r, column=2, value=f"={GAP}!B{g}")
    ws.cell(row=r, column=3, value=f"={GAP}!E{g}")
    ws.cell(row=r, column=4, value=f"={GAP}!F{g}")
    ws.cell(row=r, column=5, value=f"=COUNTIF({PR('O')},A{r})")
    ws.cell(row=r, column=6, value=f"=E{r}/D{r}").number_format = "0.0"
    ws.cell(row=r, column=7, value=f"={GAP}!Q{g}/{GAP}!N{g}").number_format = PCT1
    ws.cell(row=r, column=8, value=f"=F{r}*G{r}").number_format = "0.00"
    ws.cell(row=r, column=9, value=f"={GAP}!J{g}").number_format = BRL
    ws.cell(row=r, column=10, value=f"={GAP}!I{g}").number_format = PCT
    ws.cell(row=r, column=11, value=f"={GAP}!L{g}").number_format = BRL
    for col, base in ((12, "H"), (13, "F"), (14, "G")):
        ws.cell(row=r, column=col, value=f"=RANK({base}{r},${base}$5:${base}${E1})")
    ws.cell(row=r, column=15, value=f'=AVERAGEIFS({ex("E")},{ex("B")},A{r})').number_format = "0.0"
ws.conditional_formatting.add(f"H5:H{E1}", ColorScaleRule(start_type="min", start_color="FFFFFF",
                                                          end_type="max", end_color="63BE7B"))
rr = E1 + 2
idx = f"MATCH(1,$L$5:$L${E1},0)"
ws.cell(row=rr, column=1, value="Recomendação").font = F_SUB
ws.cell(row=rr + 1, column=1, value="Franquia recomendada (ranking 1)")
ws.cell(row=rr + 1, column=4, value=f"=INDEX($B$5:$B${E1},{idx})").fill = FILL_DESTAQUE
ws.cell(row=rr + 2, column=1, value="2ª opção")
ws.cell(row=rr + 2, column=4, value=f"=INDEX($B$5:$B${E1},MATCH(2,$L$5:$L${E1},0))")
ws.cell(row=rr + 3, column=1, value="Por quê (franquia recomendada)")
ws.cell(row=rr + 3, column=4, value=(
    f'=FIXED(INDEX($F$5:$F${E1},{idx}),1)&" endereçáveis por EC (rede: "&FIXED(SUM($E$5:$E${E1})/SUM($D$5:$D${E1}),1)&"), '
    f'conversão do funil de "&FIXED(INDEX($G$5:$G${E1},{idx})*100,1)&"% (rede: "&FIXED({GAP}!Q{F1 + 1}/{GAP}!N{F1 + 1}*100,1)&'
    f'"%), ECs com "&FIXED(INDEX($O$5:$O${E1},{idx}),1)&" leads trabalhados por mês (rede: "&FIXED(AVERAGE({ex("E")}),1)&'
    f'") e gap de R$ "&FIXED(INDEX($I$5:$I${E1},{idx}),0)'))
ws.cell(row=rr + 4, column=1, value="Sensibilidade")
ws.cell(row=rr + 4, column=4, value=(
    f'="Só pela demanda: "&INDEX($M$5:$M${E1},{idx})&"º · só pela conversão: "&INDEX($N$5:$N${E1},{idx})&"º"'))
ws.cell(row=rr + 5, column=1, value="Estados sem franquia")
ws.cell(row=rr + 5, column=4, value=(
    f'=COUNTIF({PR("O")},"Sem franquia")&" endereçáveis em MT, MA, RO e MS: sem EC nem histórico de execução para avaliar"'))
ws.cell(row=rr + 7, column=1, value="Leitura").font = F_SUB
for k, t in enumerate([
    "• O gap está nas franquias pequenas (aba Gap_Meta). O EC adicional vai para onde há mais demanda por EC, um funil que "
    "já converte e ECs já carregados.",
    "• BA e MG são as únicas entre as 3 primeiras em demanda por EC e em conversão do funil. A BA tem mais demanda por EC, "
    "os ECs mais carregados da rede (18,7 leads trabalhados por mês, contra 16,5) e está abaixo da meta (gap de R$ 3 mil); "
    "a MG já entrega 114% da meta. A margem é pequena (0,56 contra 0,50 fechamentos potenciais por EC): a diferença vem da "
    "demanda por EC, não da conversão (4,5% × 4,7%). A SP Interior tem a maior demanda por EC, mas divide com o PA a menor "
    "conversão da rede (2,9%).",
]):
    ws.cell(row=rr + 8 + k, column=1, value=t)
larguras(ws, {"A": 9, "B": 21, "C": 11, "D": 6, "E": 13, "F": 12, "G": 13, "H": 14, "I": 11, "J": 11,
              "K": 12, "L": 14, "M": 14, "N": 12, "O": 14})
ws.row_dimensions[4].height = 45

wb.move_sheet("Execucao", offset=len(wb.sheetnames))
wb.save(OUT)
print("Gerado:", OUT)
