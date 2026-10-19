"""Estilos e helpers compartilhados pelas planilhas geradas (Desafios 2 e 3)."""
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

AZUL = "1F3A5F"
F_TITULO = Font(bold=True, size=14, color=AZUL)
F_SUB = Font(bold=True, size=11, color=AZUL)
F_HEADER = Font(bold=True, color="FFFFFF")
F_ALERTA = Font(bold=True, color="C00000")
FILL_HEADER = PatternFill("solid", fgColor=AZUL)
FILL_INPUT = PatternFill("solid", fgColor="FFF2CC")
FILL_DESTAQUE = PatternFill("solid", fgColor="E2EFDA")
FILL_ALERTA = PatternFill("solid", fgColor="FCE4D6")
BORDA = Border(bottom=Side(style="thin", color="BFBFBF"))
WRAP = Alignment(wrap_text=True, vertical="top")
PCT, PCT1, NUM2, NUM4, BRL = "0%", "0.0%", "0.00", "0.0000", 'R$ #,##0'


def header(ws, row, cols, start_col=1):
    for i, txt in enumerate(cols):
        c = ws.cell(row=row, column=start_col + i, value=txt)
        c.font, c.fill, c.alignment = F_HEADER, FILL_HEADER, Alignment(wrap_text=True, vertical="center")


def larguras(ws, mapa):
    for col, w in mapa.items():
        ws.column_dimensions[col].width = w


def py(v):
    """Converte escalares numpy/pandas para tipos nativos (openpyxl não aceita numpy)."""
    return v.item() if hasattr(v, "item") else v
